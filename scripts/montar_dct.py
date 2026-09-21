# -*- coding: utf-8 -*-
"""Extrai o organizador curricular dos PDFs do DCT/TO e gera banco + planilha."""

from __future__ import annotations

import json
import re
import sqlite3
import statistics
from collections import defaultdict
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SAIDA = ROOT / "saida"

PDFS = [
    ROOT / "dct linguagens.pdf",
    ROOT / "dct ciencias da natureza e matematica.pdf",
    ROOT / "dct ciencias humanas e ensino religioso.pdf",
]

COMPONENTES = {
    "LP": ("Linguagens", "Língua Portuguesa"),
    "LI": ("Linguagens", "Língua Inglesa"),
    "AR": ("Linguagens", "Arte"),
    "EF": ("Linguagens", "Educação Física"),
    "CI": ("Ciências da Natureza", "Ciências"),
    "MA": ("Matemática", "Matemática"),
    "GE": ("Ciências Humanas", "Geografia"),
    "HI": ("Ciências Humanas", "História"),
    "ER": ("Ensino Religioso", "Ensino Religioso"),
}

COMP_PATTERNS = [
    (re.compile(r"L[ÍI]NGUA\s+PORTUGUESA", re.I), "LP"),
    (re.compile(r"L[ÍI]NGUA\s+INGLESA", re.I), "LI"),
    (re.compile(r"EDUCA[ÇC][ÃA]O\s+F[ÍI]SICA", re.I), "EF"),
    (re.compile(r"\bARTE\b", re.I), "AR"),
    (re.compile(r"CI[ÊE]NCIAS(?:\s+DA\s+NATUREZA)?", re.I), "CI"),
    (re.compile(r"MATEM[ÁA]TICA", re.I), "MA"),
    (re.compile(r"GEOGRAFIA", re.I), "GE"),
    (re.compile(r"HIST[ÓO]RIA", re.I), "HI"),
    (re.compile(r"ENSINO\s+RELIGIOSO", re.I), "ER"),
]

HEADER_RE = re.compile(
    r"(L[ÍI]NGUA\s+PORTUGUESA|L[ÍI]NGUA\s+INGLESA|EDUCA[ÇC][ÃA]O\s+F[ÍI]SICA|\bARTE\b|"
    r"CI[ÊE]NCIAS(?:\s+DA\s+NATUREZA)?|MATEM[ÁA]TICA|GEOGRAFIA|HIST[ÓO]RIA|ENSINO\s+RELIGIOSO)"
    r"[\s\-–—:]*"
    r"(\d{1,2})\s*[º°oª]?\s*(?:ANO)?"
    r"(?:.*?(\d{1,2})\s*[º°oª]?\s*BIMESTRE)?",
    re.I,
)
EIXO_LINE_RE = re.compile(r"EIXO\s*:\s*(.+)", re.I)
CODE_TOKEN_RE = re.compile(r"^\(?EF\d{2}[A-Z]{2}\d{2,3}[a-z]?(?:TO)?\)?$")
CODE_IN_TEXT_RE = re.compile(r"EF\d{2}[A-Z]{2}\d{2,3}[a-z]?(?:TO)?")
SOFT_HYPHEN = "\u00ad"

NOISE = re.compile(
    r"^(documento curricular|tocantins|organizador curricular|ensino fundamental|"
    r"anos iniciais|anos finais|unidades?|tem[áa]ticas?|habilidades?|"
    r"objetos?|conhecimento|sugest[õo]es?|pedag[óo]gicas?|metodol[óo]gicas?|"
    r"ca$|eixo$)$",
    re.I,
)

FAIXAS_ANO = {
    "12": [1, 2],
    "15": [1, 2, 3, 4, 5],
    "35": [3, 4, 5],
    "67": [6, 7],
    "68": [6, 7, 8],
    "69": [6, 7, 8, 9],
    "89": [8, 9],
}


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace(SOFT_HYPHEN, "")
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def join_words(words: list) -> str:
    parts: list[str] = []
    for w in words:
        token = w[4].replace(SOFT_HYPHEN, "")
        if not token:
            continue
        if parts and (parts[-1].endswith("-") or parts[-1].endswith(SOFT_HYPHEN)):
            parts[-1] = parts[-1].rstrip("-").rstrip(SOFT_HYPHEN) + token
        elif token in {",", ".", ";", ":", ")", "]"}:
            parts[-1:] = [((parts[-1] if parts else "") + token)]
        elif parts and parts[-1].endswith("("):
            parts[-1] += token
        else:
            parts.append(token)
    text = " ".join(parts)
    text = text.replace("( ", "(").replace(" )", ")")
    text = re.sub(r"\s+([,.;:])", r"\1", text)
    return clean_text(text)


def median_or(values: list[float], default: float) -> float:
    return statistics.median(values) if values else default


def componente_from_text(text: str) -> str | None:
    for rx, codigo in COMP_PATTERNS:
        if rx.search(text):
            return codigo
    return None


def anos_do_codigo(codigo: str) -> list[int]:
    par = codigo[2:4]
    if par in FAIXAS_ANO:
        return FAIXAS_ANO[par]
    try:
        return [int(par)]
    except ValueError:
        return []


def origem_do_codigo(codigo: str) -> str:
    return "DCT_TO" if "TO" in codigo else "BNCC"


def sigla_codigo(codigo: str) -> str:
    match = CODE_IN_TEXT_RE.search(codigo or "")
    if not match:
        return ""
    return match.group(0)[4:6]


def clean_label(text: str) -> str:
    text = clean_text(text)
    text = re.sub(r"\d+\s*Documento Curricular.*", "", text, flags=re.I)
    text = re.sub(r"Documento Curricular.*", "", text, flags=re.I)
    text = re.sub(r"\|\s*Tocantins.*", "", text, flags=re.I)
    text = re.sub(
        r"\b(UNIDADES?|TEM[ÁA]TICAS?|HABILIDADES?|OBJETOS?(?:\s+DE\s+CONHECIMENTO)?|SUGEST[ÕO]ES?|PEDAG[ÓO]GICAS?)\b",
        " ",
        text,
        flags=re.I,
    )
    text = re.sub(r"\bCAMPOS?\s+DE\s+ATUA[ÇC][ÃA]O\b", " ", text, flags=re.I)
    return clean_text(re.sub(r"\s+", " ", text))


def page_words(page) -> list:
    words = list(page.get_text("words"))
    h = page.rect.height
    kept = []
    for w in words:
        x0, y0, x1, y1, token, *_ = w
        if y0 > h - 70:
            continue
        low = token.lower()
        if "documento curricular" in low:
            continue
        kept.append((x0, y0, x1, y1, token))
    return kept


def group_lines(words: list, tol: float = 5.0) -> list[list]:
    lines: list[list] = []
    for w in sorted(words, key=lambda item: (round(item[1], 1), item[0])):
        if lines and abs(w[1] - lines[-1][0][1]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return lines


def parse_header_line(text: str) -> dict | None:
    compact = re.sub(r"\s+", " ", text)
    match = HEADER_RE.search(compact)
    if not match:
        return None
    nome, ano, bimestre = match.group(1), match.group(2), match.group(3)
    codigo = componente_from_text(nome)
    if not codigo or not ano:
        return None
    ano_n = int(ano)
    if ano_n < 1 or ano_n > 9:
        return None
    bim = int(bimestre) if bimestre else None
    if bim is not None and bim not in (1, 2, 3, 4):
        bim = None
    return {
        "componente": codigo,
        "ano": ano_n,
        "bimestre": bim,
        "texto": compact,
    }


def find_headers(words: list) -> list[dict]:
    headers = []
    for line in group_lines(words, tol=6):
        text = join_words(sorted(line, key=lambda item: item[0]))
        parsed = parse_header_line(text)
        if parsed:
            parsed["y"] = min(w[1] for w in line)
            headers.append(parsed)
    headers.sort(key=lambda item: item["y"])
    return headers


def find_eixo_line(words: list, y0: float, y1: float) -> str:
    for line in group_lines(words, tol=6):
        y = min(w[1] for w in line)
        if y < y0 or y > y1:
            continue
        text = join_words(sorted(line, key=lambda item: item[0]))
        match = EIXO_LINE_RE.search(text)
        if match:
            return clean_text(match.group(1))
    return ""


def code_words(words: list, componente: str | None = None) -> list:
    found = []
    for w in words:
        match = CODE_IN_TEXT_RE.search(w[4])
        if not match:
            continue
        if componente and sigla_codigo(match.group(0)) != componente:
            continue
        found.append(w)
    if not found:
        return []
    left = min(w[0] for w in found)
    return [w for w in found if w[0] <= left + 140]


def infer_cuts(words: list, codes: list) -> tuple[float, float, float]:
    hab_x = median_or([w[0] for w in codes], 160)
    sug_headers = [w[0] for w in words if re.search(r"SUGEST", w[4], re.I) and w[0] > hab_x + 80]
    sug_x = median_or(sug_headers, 0)
    if sug_x < hab_x + 120:
        right_xs = sorted(w[0] for w in words if w[0] > hab_x + 150)
        sug_x = largest_gap(right_xs, fallback=480)
    obj_candidates = [
        w[0]
        for w in words
        if hab_x + 145 <= w[0] < (sug_x - 20 if sug_x > hab_x + 200 else hab_x + 250)
    ]
    obj_x = min(obj_candidates) if obj_candidates else hab_x + 150
    if obj_x >= sug_x - 30 and sug_x > hab_x + 200:
        obj_x = hab_x + max(120, (sug_x - hab_x) * 0.45)
    return hab_x, obj_x, sug_x


def largest_gap(xs: list[float], fallback: float) -> float:
    if len(xs) < 8:
        return fallback
    best_gap = 0
    best_at = fallback
    for a, b in zip(xs, xs[1:]):
        gap = b - a
        if gap > best_gap and 350 < ((a + b) / 2) < 700:
            best_gap = gap
            best_at = (a + b) / 2
    return best_at if best_gap >= 20 else fallback


def column_of(x0: float, hab_x: float, obj_x: float, sug_x: float) -> str:
    if x0 < hab_x - 12:
        return "esquerda"
    if x0 < obj_x:
        return "hab"
    if x0 < sug_x - 8:
        return "obj"
    return "sug"


def is_noise(text: str) -> bool:
    compact = re.sub(r"\s+", " ", text).strip(" :-")
    if not compact:
        return True
    if NOISE.match(compact):
        return True
    if compact.isdigit() and len(compact) <= 3:
        return True
    return False


def split_unidades(text: str) -> tuple[str, str]:
    """Separa tema transversal (CTS) da unidade BNCC, quando ambos vêm à esquerda."""
    text = clean_label(text)
    cts = re.search(
        r"CI[ÊE]NCIA,?\s+TECNOLOGIA\s+E\s+SOCIEDADE",
        text,
        re.I,
    )
    tema = ""
    if cts:
        tema = clean_text(cts.group(0))
        text = clean_label(text.replace(cts.group(0), " "))
    return text, tema


def reconstruct_left(words: list) -> tuple[str, str]:
    if not words:
        return "", ""
    clusters: dict[int, list] = defaultdict(list)
    for w in words:
        if re.search(r"documento|curricular|tocantins", w[4], re.I):
            continue
        clusters[int(w[0] // 18) * 18].append(w)
    labels = []
    for x in sorted(clusters):
        group = sorted(clusters[x], key=lambda item: (item[1], item[0]))
        label = clean_label(join_words(group))
        if label and not is_noise(label):
            labels.append(label)
    return split_unidades(" ".join(labels))


def split_obj_sug(objetos: str, sugestao: str) -> tuple[str, str]:
    objetos = clean_text(objetos)
    sugestao = clean_text(sugestao)
    if sugestao or not objetos:
        return objetos, sugestao
    parts = re.split(r"(?<=[.;:])\s*-\s+", objetos, maxsplit=1)
    if len(parts) == 2 and len(parts[0]) < 450:
        return clean_text(parts[0]), "- " + clean_text(parts[1])
    parts = re.split(
        r"(?<=\.)\s+(?=O (?:professor|objetivo|desenvolvimento|trabalho)|Essa habilidade|Esta habilidade|Trata-se|Para o)",
        objetos,
        maxsplit=1,
    )
    if len(parts) == 2:
        return clean_text(parts[0]), clean_text(parts[1])
    return objetos, sugestao


def extract_codes_in_order(text: str) -> list[str]:
    codes = []
    for match in CODE_IN_TEXT_RE.finditer(text):
        codigo = match.group(0)
        if codigo not in codes:
            codes.append(codigo)
    return codes


def split_enunciados(hab_text: str) -> list[tuple[str, str]]:
    hab_text = clean_text(hab_text)
    matches = list(CODE_IN_TEXT_RE.finditer(hab_text))
    if not matches:
        return []
    rows = []
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(hab_text)
        enunciado = clean_text(hab_text[start:end].lstrip(") :.-"))
        rows.append((match.group(0), enunciado))
    return rows


def extract_page(page, pdf_name: str, page_no: int, context: dict) -> list[dict]:
    words = page_words(page)
    if not words:
        return []
    headers = find_headers(words)
    codes_all = code_words(words)
    if not codes_all and not headers:
        return []

    regions = []
    if headers:
        for i, header in enumerate(headers):
            y0 = header["y"]
            y1 = headers[i + 1]["y"] - 4 if i + 1 < len(headers) else page.rect.height - 70
            regions.append((y0, y1, header))
    elif codes_all and context.get("componente"):
        regions.append((min(w[1] for w in codes_all) - 30, page.rect.height - 70, dict(context)))

    records = []
    for y0, y1, header in regions:
        region_words = [w for w in words if y0 - 2 <= w[1] < y1]
        if not header.get("componente") and context.get("componente"):
            header = dict(context)

        context.update(
            {
                "componente": header.get("componente") or context.get("componente"),
                "ano": header.get("ano") if header.get("ano") else context.get("ano"),
                "bimestre": header.get("bimestre") if "bimestre" in header else context.get("bimestre"),
            }
        )
        componente = context.get("componente")
        eixo_extra = find_eixo_line(region_words, y0, y0 + 80)

        codes = code_words(region_words, componente)
        if not codes:
            continue
        hab_x, obj_x, sug_x = infer_cuts(region_words, codes)
        codes_sorted = sorted(codes, key=lambda w: (w[1], w[0]))

        # Agrupa códigos da mesma linha (mesma célula da tabela).
        groups: list[list] = []
        for code in codes_sorted:
            if groups and abs(code[1] - groups[-1][0][1]) <= 10:
                groups[-1].append(code)
            else:
                groups.append([code])

        for gi, group in enumerate(groups):
            row_y0 = min(w[1] for w in group) - 2
            if gi + 1 < len(groups):
                row_y1 = min(w[1] for w in groups[gi + 1]) - 2
            else:
                row_y1 = y1
            # evita engolir o próximo cabeçalho
            row_y1 = min(row_y1, y1)

            row_words = [w for w in region_words if row_y0 <= w[1] < row_y1]
            buckets = defaultdict(list)
            for w in row_words:
                buckets[column_of(w[0], hab_x, obj_x, sug_x)].append(w)

            hab_text = join_words(sorted(buckets["hab"], key=lambda item: (item[1], item[0])))
            obj_text = join_words(sorted(buckets["obj"], key=lambda item: (item[1], item[0])))
            sug_text = join_words(sorted(buckets["sug"], key=lambda item: (item[1], item[0])))
            unidade, tema = reconstruct_left(buckets["esquerda"])

            obj_text, sug_text = split_obj_sug(obj_text, sug_text)
            obj_text = clean_label(obj_text)
            sug_text = clean_label(sug_text)

            pares = [
                (codigo, enunciado)
                for codigo, enunciado in split_enunciados(hab_text)
                if sigla_codigo(codigo) == componente
            ]
            if not pares:
                for cw in group:
                    match = CODE_IN_TEXT_RE.search(cw[4])
                    if match and sigla_codigo(match.group(0)) == componente:
                        pares.append((match.group(0), ""))
            extras = [
                c
                for c in extract_codes_in_order(hab_text + " " + sug_text + " " + obj_text)
                if sigla_codigo(c) != componente
            ]

            for codigo, enunciado in pares:
                enunciado = clean_label(enunciado)
                if is_noise(enunciado) and len(enunciado) < 12:
                    enunciado = ""
                articulacoes = [
                    c
                    for c in extract_codes_in_order(enunciado + " " + sug_text + " " + obj_text)
                    if c != codigo
                ]
                for extra in extras:
                    if extra not in articulacoes:
                        articulacoes.append(extra)
                records.append(
                    {
                        "componente": componente,
                        "ano": context.get("ano"),
                        "bimestre": context.get("bimestre"),
                        "codigo": codigo,
                        "origem": origem_do_codigo(codigo),
                        "anos_codigo": anos_do_codigo(codigo),
                        "unidade_tematica": unidade,
                        "tema_transversal": tema,
                        "eixo": eixo_extra,
                        "enunciado": enunciado,
                        "objetos_conhecimento": obj_text,
                        "sugestao_pedagogica": sug_text,
                        "articulacoes": articulacoes,
                        "fonte_pdf": pdf_name,
                        "pagina": page_no,
                    }
                )
    return records


def fill_forward(rows: list[dict]) -> list[dict]:
    last = None
    for row in rows:
        same = (
            last
            and last.get("componente") == row.get("componente")
            and last.get("ano") == row.get("ano")
            and last.get("bimestre") == row.get("bimestre")
            and last.get("fonte_pdf") == row.get("fonte_pdf")
        )
        if same:
            if not row.get("unidade_tematica"):
                row["unidade_tematica"] = last.get("unidade_tematica") or ""
            if not row.get("tema_transversal"):
                row["tema_transversal"] = last.get("tema_transversal") or ""
            if not row.get("eixo"):
                row["eixo"] = last.get("eixo") or ""
        last = row
    return rows


def extract_all() -> list[dict]:
    rows: list[dict] = []
    for pdf_path in PDFS:
        doc = pymupdf.open(pdf_path)
        context: dict = {}
        print(f"Extraindo {pdf_path.name} ({doc.page_count} páginas)...")
        for i, page in enumerate(doc):
            page_rows = extract_page(page, pdf_path.name, i + 1, context)
            if page_rows:
                rows.extend(page_rows)
        doc.close()
    rows = [r for r in rows if r.get("componente") and r.get("codigo")]
    rows = [r for r in rows if r.get("ano") is None or r.get("ano") in range(1, 10)]
    return fill_forward(rows)


def merge_best_enunciado(rows: list[dict]) -> list[dict]:
    best: dict[str, str] = {}
    for row in rows:
        key = f"{row['componente']}|{row['codigo']}"
        text = row.get("enunciado") or ""
        if len(text) > len(best.get(key, "")):
            best[key] = text
    for row in rows:
        key = f"{row['componente']}|{row['codigo']}"
        if (not row.get("enunciado") or len(row["enunciado"]) < 40) and best.get(key):
            if len(best[key]) > len(row.get("enunciado") or ""):
                row["enunciado"] = best[key]
    return rows


def build_sqlite(rows: list[dict], path: Path) -> None:
    if path.exists():
        path.unlink()
    conn = sqlite3.connect(path)
    conn.execute("PRAGMA foreign_keys = ON")
    cur = conn.cursor()
    cur.executescript(
        """
        CREATE TABLE area (
            id INTEGER PRIMARY KEY,
            nome TEXT NOT NULL UNIQUE
        );
        CREATE TABLE componente (
            id INTEGER PRIMARY KEY,
            area_id INTEGER NOT NULL REFERENCES area(id),
            nome TEXT NOT NULL,
            codigo TEXT NOT NULL UNIQUE
        );
        CREATE TABLE habilidade (
            id INTEGER PRIMARY KEY,
            componente_id INTEGER NOT NULL REFERENCES componente(id),
            codigo TEXT NOT NULL,
            enunciado TEXT,
            origem TEXT NOT NULL,
            anos_do_codigo TEXT,
            UNIQUE (componente_id, codigo)
        );
        CREATE TABLE alocacao (
            id INTEGER PRIMARY KEY,
            habilidade_id INTEGER NOT NULL REFERENCES habilidade(id),
            ano INTEGER,
            bimestre INTEGER,
            unidade_tematica TEXT,
            eixo TEXT,
            tema_transversal TEXT,
            objetos_conhecimento TEXT,
            sugestao_pedagogica TEXT,
            fonte_pdf TEXT,
            pagina INTEGER
        );
        CREATE TABLE articulacao (
            id INTEGER PRIMARY KEY,
            de_codigo TEXT NOT NULL,
            para_codigo TEXT NOT NULL,
            UNIQUE (de_codigo, para_codigo)
        );
        CREATE VIEW v_consulta AS
        SELECT
            a.nome AS area,
            c.nome AS componente,
            c.codigo AS componente_codigo,
            al.ano,
            al.bimestre,
            al.unidade_tematica,
            al.eixo,
            al.tema_transversal,
            h.codigo,
            h.origem,
            h.enunciado,
            al.objetos_conhecimento,
            al.sugestao_pedagogica,
            h.anos_do_codigo,
            al.fonte_pdf,
            al.pagina
        FROM alocacao al
        JOIN habilidade h ON h.id = al.habilidade_id
        JOIN componente c ON c.id = h.componente_id
        JOIN area a ON a.id = c.area_id
        ORDER BY c.codigo, al.ano, al.bimestre, h.codigo;
        """
    )

    area_ids = {}
    comp_ids = {}
    for codigo, (area, nome) in COMPONENTES.items():
        if area not in area_ids:
            cur.execute("INSERT INTO area(nome) VALUES (?)", (area,))
            area_ids[area] = cur.lastrowid
        cur.execute(
            "INSERT INTO componente(area_id, nome, codigo) VALUES (?, ?, ?)",
            (area_ids[area], nome, codigo),
        )
        comp_ids[codigo] = cur.lastrowid

    hab_ids = {}
    for row in rows:
        key = (row["componente"], row["codigo"])
        if key in hab_ids:
            continue
        cur.execute(
            """
            INSERT INTO habilidade(componente_id, codigo, enunciado, origem, anos_do_codigo)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                comp_ids[row["componente"]],
                row["codigo"],
                row.get("enunciado") or "",
                row["origem"],
                ",".join(str(n) for n in row.get("anos_codigo") or []),
            ),
        )
        hab_ids[key] = cur.lastrowid

    seen_aloc = set()
    for row in rows:
        key = (
            row["componente"],
            row["codigo"],
            row.get("ano"),
            row.get("bimestre"),
            row.get("unidade_tematica") or "",
        )
        if key in seen_aloc:
            continue
        seen_aloc.add(key)
        cur.execute(
            """
            INSERT INTO alocacao(
                habilidade_id, ano, bimestre, unidade_tematica, eixo, tema_transversal,
                objetos_conhecimento, sugestao_pedagogica, fonte_pdf, pagina
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                hab_ids[(row["componente"], row["codigo"])],
                row.get("ano"),
                row.get("bimestre"),
                row.get("unidade_tematica") or "",
                row.get("eixo") or "",
                row.get("tema_transversal") or "",
                row.get("objetos_conhecimento") or "",
                row.get("sugestao_pedagogica") or "",
                row.get("fonte_pdf"),
                row.get("pagina"),
            ),
        )

    seen_art = set()
    for row in rows:
        for dest in row.get("articulacoes") or []:
            pair = (row["codigo"], dest)
            if pair in seen_art:
                continue
            seen_art.add(pair)
            cur.execute(
                "INSERT OR IGNORE INTO articulacao(de_codigo, para_codigo) VALUES (?, ?)",
                pair,
            )

    conn.commit()
    conn.close()


def build_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")

    columns = [
        "area",
        "componente",
        "codigo",
        "origem",
        "ano",
        "bimestre",
        "unidade_tematica",
        "eixo",
        "tema_transversal",
        "enunciado",
        "objetos_conhecimento",
        "sugestao_pedagogica",
        "articulacoes",
        "anos_do_codigo",
        "fonte_pdf",
        "pagina",
    ]

    def write_sheet(ws, data: list[dict]) -> None:
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in data:
            area, nome = COMPONENTES[row["componente"]]
            ws.append(
                [
                    area,
                    nome,
                    row["codigo"],
                    row["origem"],
                    row.get("ano"),
                    row.get("bimestre"),
                    row.get("unidade_tematica") or "",
                    row.get("eixo") or "",
                    row.get("tema_transversal") or "",
                    row.get("enunciado") or "",
                    row.get("objetos_conhecimento") or "",
                    row.get("sugestao_pedagogica") or "",
                    ", ".join(row.get("articulacoes") or []),
                    ",".join(str(n) for n in row.get("anos_codigo") or []),
                    row.get("fonte_pdf") or "",
                    row.get("pagina"),
                ]
            )
        for col in range(1, len(columns) + 1):
            letter = get_column_letter(col)
            width = 18
            if columns[col - 1] in {"enunciado", "objetos_conhecimento", "sugestao_pedagogica"}:
                width = 48
            if columns[col - 1] in {"unidade_tematica", "eixo"}:
                width = 28
            ws.column_dimensions[letter].width = width
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(columns)):
            for cell in row:
                cell.alignment = wrap
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 22

    # LEIA-ME
    ws = wb.active
    ws.title = "LEIA-ME"
    ws["A1"] = "DCT Tocantins — organizador curricular (Ensino Fundamental)"
    ws["A1"].font = Font(bold=True, size=14)
    notes = [
        "",
        "Esta planilha é a camada de conferência. O arquivo dct_tocantins.sqlite é o formato para o sistema.",
        "Cada linha é uma habilidade colocada em um ano/bimestre, como no quadro do PDF.",
        "",
        "Colunas:",
        "codigo — código BNCC (EF06CI02) ou código próprio do Tocantins (sufixo TO).",
        "origem — BNCC ou DCT_TO.",
        "ano / bimestre — recorte do organizador curricular do DCT.",
        "unidade_tematica — unidade BNCC ou campo de atuação (em Língua Portuguesa).",
        "eixo — eixo de linguagem, quando o quadro trouxer (LP/Inglês).",
        "tema_transversal — em Ciências, 'Ciência, Tecnologia e Sociedade'.",
        "enunciado — texto da habilidade.",
        "objetos_conhecimento / sugestao_pedagogica — colunas do quadro.",
        "articulacoes — outros códigos citados na mesma célula.",
        "fonte_pdf / pagina — de onde foi extraído, para conferência.",
        "",
        "Há uma aba por componente e a aba TODAS com o conjunto completo.",
        "Textos introdutórios dos cadernos não entram aqui; ficam nos PDFs originais.",
        "A extração é automática a partir do layout dos PDFs. Vale conferir amostra por componente.",
    ]
    for i, line in enumerate(notes, start=2):
        ws[f"A{i}"] = line
    ws.column_dimensions["A"].width = 110

    by_comp: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_comp[row["componente"]].append(row)

    ordem = ["LP", "LI", "AR", "EF", "CI", "MA", "GE", "HI", "ER"]
    for codigo in ordem:
        data = by_comp.get(codigo, [])
        data.sort(key=lambda r: (r.get("ano") or 0, r.get("bimestre") or 0, r["codigo"], r.get("pagina") or 0))
        ws_comp = wb.create_sheet(COMPONENTES[codigo][1][:31])
        write_sheet(ws_comp, data)

    todas = sorted(
        rows,
        key=lambda r: (r["componente"], r.get("ano") or 0, r.get("bimestre") or 0, r["codigo"]),
    )
    write_sheet(wb.create_sheet("TODAS"), todas)
    wb.save(path)


def report(rows: list[dict]) -> dict:
    by_comp = defaultdict(lambda: {"alocacoes": 0, "codigos": set(), "to": 0, "sem_enunciado": 0})
    for row in rows:
        info = by_comp[row["componente"]]
        info["alocacoes"] += 1
        info["codigos"].add(row["codigo"])
        if row["origem"] == "DCT_TO":
            info["to"] += 1
        if len(row.get("enunciado") or "") < 20:
            info["sem_enunciado"] += 1
    resumo = {
        codigo: {
            "nome": COMPONENTES[codigo][1],
            "alocacoes": info["alocacoes"],
            "habilidades": len(info["codigos"]),
            "codigos_TO": info["to"],
            "sem_enunciado": info["sem_enunciado"],
        }
        for codigo, info in sorted(by_comp.items())
    }
    return {
        "alocacoes": len(rows),
        "habilidades_unicas": len({(r["componente"], r["codigo"]) for r in rows}),
        "componentes": resumo,
    }


def main() -> None:
    SAIDA.mkdir(exist_ok=True)
    rows = extract_all()
    rows = merge_best_enunciado(rows)
    resumo = report(rows)
    print(json.dumps(resumo, ensure_ascii=False, indent=2))

    sqlite_path = SAIDA / "dct_tocantins.sqlite"
    xlsx_path = SAIDA / "dct_tocantins.xlsx"
    json_path = SAIDA / "dct_tocantins.json"
    report_path = SAIDA / "resumo_extracao.json"

    build_sqlite(rows, sqlite_path)
    build_excel(rows, xlsx_path)
    json_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    report_path.write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Gerado: {sqlite_path}")
    print(f"Gerado: {xlsx_path}")
    print(f"Gerado: {json_path}")


if __name__ == "__main__":
    main()
